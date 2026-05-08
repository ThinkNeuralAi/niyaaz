"""
Daily Alert Report Generator & Email Sender
Generates an Excel report of yesterday's alerts per store and sends via email at 10:30 AM daily.
"""
import os
import logging
import smtplib
import threading
import time
from datetime import datetime, timedelta
import calendar
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email.mime.text import MIMEText
from email import encoders
from zoneinfo import ZoneInfo
from pathlib import Path

logger = logging.getLogger(__name__)

IST = ZoneInfo("Asia/Kolkata")

# All alert usecase categories and the DB tables they map to
USECASE_CONFIG = {
    'Queue & Wait Time': {
        'table': 'queue_violations',
    },
    'Uniform Compliance': {
        'table': 'dresscode_alerts',
    },
    'PPE Violations': {
        'table': 'ppe_alerts',
    },
    'Table Cleanliness': {
        'table': 'table_cleanliness_violations',
    },
    'Service Discipline': {
        'table': 'alert_gifs',
        'alert_type_filter': 'service_discipline_alert',
    },
    'Unauthorised Entry': {
        'table': 'alert_gifs',
        'alert_type_filter': 'unauthorized_entry',
    },
      'Material Theft': {
        'table': 'alert_gifs',
        'alert_type_filter': 'material_theft',
    },
    'Fall Detection': {
        'table': 'fall_snapshots',
    },
    'Smoke & Fire Detection': {
        'table': 'smoking_snapshots',
    },
    'Smoking Detection': {
        'table': 'alert_gifs',
        'alert_type_filter': 'person_smoking_detected',
    },
    'Crowd Detection': {
        'table': 'alert_gifs',
        'alert_type_filter': 'crowd',
    },
    'Idle Time': {
        'table': 'idle_time_violations',
    },
}


def _get_yesterday_range():
    """Return (start, end) datetime for yesterday in IST"""
    now = datetime.now(IST)
    yesterday = now - timedelta(days=1)
    start = yesterday.replace(hour=0, minute=0, second=0, microsecond=0)
    end = yesterday.replace(hour=23, minute=59, second=59, microsecond=999999)
    return start, end


def _daily_report_exists_for_date(report_date_str):
    """Return True if a daily report file exists for YYYY-MM-DD."""
    reports_dir = Path(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))) / 'data' / 'reports'
    if not reports_dir.exists():
        return False
    pattern = f'Daily_Alerts_Report_{report_date_str}*.xlsx'
    return any(reports_dir.glob(pattern))


def _parse_month_input(target_month):
    """Parse YYYY-MM input and return (year, month)."""
    if not target_month:
        raise ValueError("target_month is required in YYYY-MM format")

    if isinstance(target_month, str):
        parsed = datetime.strptime(target_month, '%Y-%m')
        return parsed.year, parsed.month

    # Support date/datetime-like inputs that carry year/month.
    if hasattr(target_month, 'year') and hasattr(target_month, 'month'):
        return int(target_month.year), int(target_month.month)

    raise ValueError("target_month must be a string in YYYY-MM format")


def _get_month_range(year, month):
    """Return month start/end datetimes in IST."""
    last_day = calendar.monthrange(year, month)[1]
    start = datetime(year, month, 1, 0, 0, 0, 0, tzinfo=IST)
    end = datetime(year, month, last_day, 23, 59, 59, 999999, tzinfo=IST)
    return start, end


def _get_previous_month_range(reference_dt=None):
    """Return previous calendar month range based on IST now/reference."""
    now = reference_dt or datetime.now(IST)
    first_of_this_month = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    previous_month_last_moment = first_of_this_month - timedelta(microseconds=1)
    return _get_month_range(previous_month_last_moment.year, previous_month_last_moment.month)


def _count_alerts_sql(db_manager, channel_ids, usecase_cfg, yesterday_start, yesterday_end):
    """Count alerts using raw SQL - reliable and efficient"""
    from sqlalchemy import text

    table = usecase_cfg['table']
    alert_type_filter = usecase_cfg.get('alert_type_filter')

    try:
        # Build WHERE clauses
        conditions = ["created_at >= :start_dt AND created_at <= :end_dt"]
        params = {
            'start_dt': yesterday_start.replace(tzinfo=None),
            'end_dt': yesterday_end.replace(tzinfo=None),
        }

        if channel_ids:
            placeholders = ', '.join(f':ch_{i}' for i in range(len(channel_ids)))
            conditions.append(f"channel_id IN ({placeholders})")
            for i, cid in enumerate(channel_ids):
                params[f'ch_{i}'] = cid

        if alert_type_filter:
            conditions.append("LOWER(alert_type) LIKE :atype")
            params['atype'] = f'{alert_type_filter}%'

        where = ' AND '.join(conditions)
        sql = text(f"SELECT COUNT(*) FROM {table} WHERE {where}")

        result = db_manager.db.session.execute(sql, params)
        count = result.scalar() or 0
        return count
    except Exception as e:
        try:
            db_manager.db.session.rollback()
        except Exception:
            pass
        logger.error(f"Error counting alerts in {table}: {e}")
        return 0


def generate_daily_report(app, db_manager, target_date=None, store_id=None):
    """Generate Excel report for a specific date's alerts across all stores and usecases.
    
    Args:
        app: Flask app instance
        db_manager: DatabaseManager instance
        target_date: Optional date object or string (YYYY-MM-DD). Defaults to yesterday.
        store_id: Optional store_id to filter report to a single store.
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    except ImportError:
        logger.error("openpyxl is required for daily reports. Install with: pip install openpyxl")
        return None

    if target_date:
        if isinstance(target_date, str):
            from datetime import date as date_cls
            target_date = datetime.strptime(target_date, '%Y-%m-%d').date()
        report_start = datetime.combine(target_date, datetime.min.time()).replace(tzinfo=IST)
        report_end = datetime.combine(target_date, datetime.max.time()).replace(tzinfo=IST)
    else:
        report_start, report_end = _get_yesterday_range()

    report_date = report_start.strftime('%Y-%m-%d')
    logger.info(f"Generating daily alert report for {report_date}")

    with app.app_context():
        # Get all active stores
        stores = db_manager.get_all_stores()
        active_stores = [s for s in stores if s.get('is_active')]

        # Filter to specific store if requested
        if store_id:
            active_stores = [s for s in active_stores if s.get('store_id') == store_id]

        if not active_stores:
            logger.warning("No active stores found, skipping report generation")
            return None

        # Build channel→store mapping
        all_links = db_manager.get_all_rtsp_links()
        store_channels = {}  # {store_id: [channel_ids]}
        for link in all_links:
            if link.get('is_active') and link.get('channel_id'):
                sid = link.get('store_id', '')
                store_channels.setdefault(sid, []).append(link['channel_id'])

        # Usecase column names (ordered)
        usecase_names = list(USECASE_CONFIG.keys())

        # Collect data: {store_name: {usecase: count}}
        report_data = []
        for store in active_stores:
            sid = store.get('store_id', '')
            store_name = store.get('name', sid)
            channels = store_channels.get(sid, [])

            row = {'store_name': store_name, 'usecases': {}}
            for uc_name, uc_cfg in USECASE_CONFIG.items():
                cnt = _count_alerts_sql(
                    db_manager, channels, uc_cfg,
                    report_start, report_end
                )
                row['usecases'][uc_name] = cnt
            report_data.append(row)

    # Create Excel workbook matching the sample format
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Daily AI Alerts"

    # Styles
    header_font = Font(name='Calibri', bold=True, size=14, color='FFFFFF')
    title_font = Font(name='Calibri', bold=True, size=18, color='1F4E79')
    subtitle_font = Font(name='Calibri', bold=True, size=12, color='1F4E79')
    col_header_font = Font(name='Calibri', bold=True, size=10, color='FFFFFF')
    data_font = Font(name='Calibri', size=10)
    total_font = Font(name='Calibri', bold=True, size=10, color='FFFFFF')

    header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    col_header_fill = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
    total_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    alt_row_fill = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')

    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)

    # Row 1: blank
    # Row 2: Title
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(usecase_names) + 3)
    title_cell = ws.cell(row=2, column=1, value='DAILY AI ALERTS REPORT')
    title_cell.font = title_font
    title_cell.alignment = center_align

    # Row 3: Date subtitle
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(usecase_names) + 3)
    date_cell = ws.cell(row=3, column=1, value=f'Report Date: {report_date}')
    date_cell.font = subtitle_font
    date_cell.alignment = center_align

    # Row 4: blank
    # Row 5: Column headers
    headers = ['Date', 'Location/Outlet', 'Total Alerts'] + usecase_names + ['Alerts %']
    header_row = 5
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = col_header_font
        cell.fill = col_header_fill
        cell.alignment = center_align
        cell.border = thin_border

    # Data rows
    data_start_row = 6
    grand_total = 0
    for row_idx, row_data in enumerate(report_data):
        excel_row = data_start_row + row_idx
        store_name = row_data['store_name']
        usecase_counts = [row_data['usecases'].get(uc, 0) for uc in usecase_names]
        row_total = sum(usecase_counts)
        grand_total += row_total

        # Date
        cell = ws.cell(row=excel_row, column=1, value=report_start.strftime('%Y-%m-%d'))
        cell.font = data_font
        cell.alignment = center_align
        cell.border = thin_border

        # Store name
        cell = ws.cell(row=excel_row, column=2, value=store_name)
        cell.font = data_font
        cell.alignment = left_align
        cell.border = thin_border

        # Total (formula)
        last_uc_col = 3 + len(usecase_names)
        from openpyxl.utils import get_column_letter
        total_formula = f'=SUM({get_column_letter(4)}{excel_row}:{get_column_letter(last_uc_col)}{excel_row})'
        cell = ws.cell(row=excel_row, column=3, value=total_formula)
        cell.font = Font(name='Calibri', bold=True, size=10)
        cell.alignment = center_align
        cell.border = thin_border

        # Usecase counts
        for uc_idx, count in enumerate(usecase_counts):
            cell = ws.cell(row=excel_row, column=4 + uc_idx, value=count)
            cell.font = data_font
            cell.alignment = center_align
            cell.border = thin_border

        # Alerts % (will use formula referencing total row)
        cell = ws.cell(row=excel_row, column=last_uc_col + 1)
        cell.font = data_font
        cell.alignment = center_align
        cell.border = thin_border

        # Alternate row colors
        if row_idx % 2 == 1:
            for c in range(1, last_uc_col + 2):
                ws.cell(row=excel_row, column=c).fill = alt_row_fill

    # Total row
    total_row = data_start_row + len(report_data)
    last_uc_col = 3 + len(usecase_names)

    cell = ws.cell(row=total_row, column=1)
    cell.border = thin_border

    cell = ws.cell(row=total_row, column=2, value='Total')
    cell.font = total_font
    cell.fill = total_fill
    cell.alignment = center_align
    cell.border = thin_border

    # Grand total formula
    from openpyxl.utils import get_column_letter
    total_formula = f'=SUM(C{data_start_row}:C{total_row - 1})'
    cell = ws.cell(row=total_row, column=3, value=total_formula)
    cell.font = total_font
    cell.fill = total_fill
    cell.alignment = center_align
    cell.border = thin_border

    # Per-usecase totals
    for uc_idx in range(len(usecase_names)):
        col = 4 + uc_idx
        col_letter = get_column_letter(col)
        formula = f'=SUM({col_letter}{data_start_row}:{col_letter}{total_row - 1})'
        cell = ws.cell(row=total_row, column=col, value=formula)
        cell.font = total_font
        cell.fill = total_fill
        cell.alignment = center_align
        cell.border = thin_border

    # Alerts % total
    cell = ws.cell(row=total_row, column=last_uc_col + 1)
    cell.font = total_font
    cell.fill = total_fill
    cell.border = thin_border

    # Now fill in % formulas (referencing grand total row)
    total_col_letter = 'C'
    total_cell_ref = f'${total_col_letter}${total_row}'
    for row_idx in range(len(report_data)):
        excel_row = data_start_row + row_idx
        formula = f'=IF({total_cell_ref}=0,0,ROUND(C{excel_row}/{total_cell_ref}*100,1))'
        ws.cell(row=excel_row, column=last_uc_col + 1, value=formula)

    # Set column widths
    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 14
    for uc_idx in range(len(usecase_names)):
        col_letter = get_column_letter(4 + uc_idx)
        ws.column_dimensions[col_letter].width = 18
    ws.column_dimensions[get_column_letter(last_uc_col + 1)].width = 12

    # Save to file
    reports_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'data', 'reports')
    os.makedirs(reports_dir, exist_ok=True)
    store_suffix = f'_{store_id}' if store_id else ''
    filename = f'Daily_Alerts_Report_{report_date}{store_suffix}.xlsx'
    filepath = os.path.join(reports_dir, filename)
    # If file is locked (e.g. open in Excel), use a timestamped name
    try:
        wb.save(filepath)
    except PermissionError:
        ts = datetime.now(IST).strftime('%H%M%S')
        filename = f'Daily_Alerts_Report_{report_date}_{ts}.xlsx'
        filepath = os.path.join(reports_dir, filename)
        wb.save(filepath)
    logger.info(f"Daily report saved to {filepath}")
    return filepath


def generate_monthly_report(app, db_manager, target_month=None, store_id=None):
    """Generate Excel report for a full month across all stores and usecases.

    Args:
        app: Flask app instance
        db_manager: DatabaseManager instance
        target_month: Optional month string (YYYY-MM). Defaults to previous month.
        store_id: Optional store_id to filter report to a single store.
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        logger.error("openpyxl is required for monthly reports. Install with: pip install openpyxl")
        return None

    if target_month:
        year, month = _parse_month_input(target_month)
        report_start, report_end = _get_month_range(year, month)
    else:
        report_start, report_end = _get_previous_month_range()

    month_key = report_start.strftime('%Y-%m')
    month_label = report_start.strftime('%B %Y')
    logger.info(f"Generating monthly alert report for {month_label}")

    with app.app_context():
        stores = db_manager.get_all_stores()
        active_stores = [s for s in stores if s.get('is_active')]

        if store_id:
            active_stores = [s for s in active_stores if s.get('store_id') == store_id]

        if not active_stores:
            logger.warning("No active stores found, skipping monthly report generation")
            return None

        all_links = db_manager.get_all_rtsp_links()
        store_channels = {}
        for link in all_links:
            if link.get('is_active') and link.get('channel_id'):
                sid = link.get('store_id', '')
                store_channels.setdefault(sid, []).append(link['channel_id'])

        usecase_names = list(USECASE_CONFIG.keys())
        report_data = []
        for store in active_stores:
            sid = store.get('store_id', '')
            store_name = store.get('name', sid)
            channels = store_channels.get(sid, [])

            row = {'store_name': store_name, 'usecases': {}}
            for uc_name, uc_cfg in USECASE_CONFIG.items():
                cnt = _count_alerts_sql(
                    db_manager, channels, uc_cfg,
                    report_start, report_end
                )
                row['usecases'][uc_name] = cnt
            report_data.append(row)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Monthly AI Alerts"

    title_font = Font(name='Calibri', bold=True, size=18, color='1F4E79')
    subtitle_font = Font(name='Calibri', bold=True, size=12, color='1F4E79')
    col_header_font = Font(name='Calibri', bold=True, size=10, color='FFFFFF')
    data_font = Font(name='Calibri', size=10)
    total_font = Font(name='Calibri', bold=True, size=10, color='FFFFFF')

    col_header_fill = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
    total_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    alt_row_fill = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')

    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)

    usecase_names = list(USECASE_CONFIG.keys())
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(usecase_names) + 3)
    title_cell = ws.cell(row=2, column=1, value='MONTHLY AI ALERTS REPORT')
    title_cell.font = title_font
    title_cell.alignment = center_align

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(usecase_names) + 3)
    date_cell = ws.cell(row=3, column=1, value=f'Report Month: {month_label}')
    date_cell.font = subtitle_font
    date_cell.alignment = center_align

    headers = ['Month', 'Location/Outlet', 'Total Alerts'] + usecase_names + ['Alerts %']
    header_row = 5
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = col_header_font
        cell.fill = col_header_fill
        cell.alignment = center_align
        cell.border = thin_border

    data_start_row = 6
    for row_idx, row_data in enumerate(report_data):
        excel_row = data_start_row + row_idx
        store_name = row_data['store_name']
        usecase_counts = [row_data['usecases'].get(uc, 0) for uc in usecase_names]

        cell = ws.cell(row=excel_row, column=1, value=month_key)
        cell.font = data_font
        cell.alignment = center_align
        cell.border = thin_border

        cell = ws.cell(row=excel_row, column=2, value=store_name)
        cell.font = data_font
        cell.alignment = left_align
        cell.border = thin_border

        last_uc_col = 3 + len(usecase_names)
        total_formula = f'=SUM({get_column_letter(4)}{excel_row}:{get_column_letter(last_uc_col)}{excel_row})'
        cell = ws.cell(row=excel_row, column=3, value=total_formula)
        cell.font = Font(name='Calibri', bold=True, size=10)
        cell.alignment = center_align
        cell.border = thin_border

        for uc_idx, count in enumerate(usecase_counts):
            cell = ws.cell(row=excel_row, column=4 + uc_idx, value=count)
            cell.font = data_font
            cell.alignment = center_align
            cell.border = thin_border

        cell = ws.cell(row=excel_row, column=last_uc_col + 1)
        cell.font = data_font
        cell.alignment = center_align
        cell.border = thin_border

        if row_idx % 2 == 1:
            for c in range(1, last_uc_col + 2):
                ws.cell(row=excel_row, column=c).fill = alt_row_fill

    total_row = data_start_row + len(report_data)
    last_uc_col = 3 + len(usecase_names)

    cell = ws.cell(row=total_row, column=1)
    cell.border = thin_border

    cell = ws.cell(row=total_row, column=2, value='Total')
    cell.font = total_font
    cell.fill = total_fill
    cell.alignment = center_align
    cell.border = thin_border

    total_formula = f'=SUM(C{data_start_row}:C{total_row - 1})'
    cell = ws.cell(row=total_row, column=3, value=total_formula)
    cell.font = total_font
    cell.fill = total_fill
    cell.alignment = center_align
    cell.border = thin_border

    for uc_idx in range(len(usecase_names)):
        col = 4 + uc_idx
        col_letter = get_column_letter(col)
        formula = f'=SUM({col_letter}{data_start_row}:{col_letter}{total_row - 1})'
        cell = ws.cell(row=total_row, column=col, value=formula)
        cell.font = total_font
        cell.fill = total_fill
        cell.alignment = center_align
        cell.border = thin_border

    cell = ws.cell(row=total_row, column=last_uc_col + 1)
    cell.font = total_font
    cell.fill = total_fill
    cell.border = thin_border

    total_col_letter = 'C'
    total_cell_ref = f'${total_col_letter}${total_row}'
    for row_idx in range(len(report_data)):
        excel_row = data_start_row + row_idx
        formula = f'=IF({total_cell_ref}=0,0,ROUND(C{excel_row}/{total_cell_ref}*100,1))'
        ws.cell(row=excel_row, column=last_uc_col + 1, value=formula)

    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 14
    for uc_idx in range(len(usecase_names)):
        col_letter = get_column_letter(4 + uc_idx)
        ws.column_dimensions[col_letter].width = 18
    ws.column_dimensions[get_column_letter(last_uc_col + 1)].width = 12

    reports_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'data', 'reports')
    os.makedirs(reports_dir, exist_ok=True)
    store_suffix = f'_{store_id}' if store_id else ''
    filename = f'Monthly_Alerts_Report_{month_key}{store_suffix}.xlsx'
    filepath = os.path.join(reports_dir, filename)
    try:
        wb.save(filepath)
    except PermissionError:
        ts = datetime.now(IST).strftime('%H%M%S')
        filename = f'Monthly_Alerts_Report_{month_key}_{ts}.xlsx'
        filepath = os.path.join(reports_dir, filename)
        wb.save(filepath)

    logger.info(f"Monthly report saved to {filepath}")
    return filepath


def send_report_email(filepath, email_config, report_subject=None, report_body=None):
    """Send the Excel report via email."""
    if not filepath or not os.path.exists(filepath):
        logger.error("Report file not found, cannot send email")
        return False

    smtp_host = email_config.get('smtp_host', '').strip()
    smtp_port = email_config.get('smtp_port', 587)
    smtp_user = email_config.get('smtp_user', '').strip()
    smtp_password = email_config.get('smtp_password', '').strip()
    sender = email_config.get('sender', smtp_user).strip()
    recipients = email_config.get('recipients', [])

    if not smtp_host or not smtp_user or not smtp_password:
        logger.error("Email configuration incomplete. Set EMAIL_SMTP_HOST, EMAIL_SMTP_USER, EMAIL_SMTP_PASSWORD env vars.")
        return False
    
    if not recipients or (isinstance(recipients, list) and len(recipients) == 0):
        logger.error("No email recipients configured. Set EMAIL_RECIPIENTS env var.")
        return False

    filename = os.path.basename(filepath)

    msg = MIMEMultipart()
    msg['From'] = sender
    msg['To'] = ', '.join(recipients)
    default_report_date = (datetime.now(IST) - timedelta(days=1)).strftime('%Y-%m-%d')
    msg['Subject'] = report_subject or f'Daily AI Alerts Report - {default_report_date}'

    body = report_body or f"""Dear Team,

Please find attached the Daily AI Alerts Report for {default_report_date}.

This is an automated report generated by Sakshi.AI Video Analytics Platform.

Regards,
Sakshi.AI System"""

    msg.attach(MIMEText(body, 'plain'))

    # Attach Excel file
    try:
        with open(filepath, 'rb') as f:
            file_bytes = f.read()
        
        part = MIMEBase('application', 'vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        part.set_payload(file_bytes)
        encoders.encode_base64(part)
        part.add_header('Content-Disposition', 'attachment', filename=filename)
        part.add_header('Content-Transfer-Encoding', 'base64')
        msg.attach(part)
        logger.info(f"Excel file attached: {filename} ({len(file_bytes)} bytes)")
    except Exception as e:
        logger.error(f"Failed to attach Excel file: {e}")
        return False

    try:
        use_ssl = email_config.get('use_ssl', False)
        logger.info(f"Sending email to {recipients} via {smtp_host}:{smtp_port} (SSL={use_ssl})")
        
        if use_ssl:
            server = smtplib.SMTP_SSL(smtp_host, smtp_port, timeout=30)
        else:
            server = smtplib.SMTP(smtp_host, smtp_port, timeout=30)
            server.starttls()
        
        server.login(smtp_user, smtp_password)
        logger.info(f"SMTP login successful for {smtp_user}")
        
        server.sendmail(sender, recipients, msg.as_string())
        server.quit()
        logger.info(f"Daily report email successfully sent to {recipients}")
        return True
    except smtplib.SMTPAuthenticationError as e:
        logger.error(f"SMTP Authentication failed: {e}")
        return False
    except smtplib.SMTPException as e:
        logger.error(f"SMTP error while sending daily report email: {e}")
        return False
    except Exception as e:
        logger.error(f"Failed to send daily report email: {e}", exc_info=True)
        return False


def get_email_config():
    """Load email configuration from environment variables or config file"""
    smtp_port_env = os.getenv('EMAIL_SMTP_PORT', '587')
    try:
        smtp_port = int(smtp_port_env)
    except (TypeError, ValueError):
        smtp_port = 587

    config = {
        'smtp_host': os.getenv('EMAIL_SMTP_HOST', ''),
        'smtp_port': smtp_port,
        'smtp_user': os.getenv('EMAIL_SMTP_USER', ''),
        'smtp_password': os.getenv('EMAIL_SMTP_PASSWORD', ''),
        'sender': os.getenv('EMAIL_SENDER', ''),
        'recipients': [],
        'use_ssl': os.getenv('EMAIL_USE_SSL', 'false').lower() == 'true',
    }

    # Recipients from env (comma-separated)
    recipients_str = os.getenv('EMAIL_RECIPIENTS', '')
    if recipients_str:
        config['recipients'] = [r.strip() for r in recipients_str.split(',') if r.strip()]

    # Fill missing values from config file. Environment variables keep precedence.
    try:
        import json
        config_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'config', 'default.json')
        with open(config_path, 'r') as f:
            file_config = json.load(f)
            email_cfg = file_config.get('email', {})

            if not config['smtp_host']:
                config['smtp_host'] = (email_cfg.get('smtp_host', '') or '').strip()

            if os.getenv('EMAIL_SMTP_PORT') is None:
                file_port = email_cfg.get('smtp_port')
                if file_port is not None:
                    try:
                        config['smtp_port'] = int(file_port)
                    except (TypeError, ValueError):
                        pass

            if not config['smtp_user']:
                config['smtp_user'] = (email_cfg.get('smtp_user', '') or '').strip()

            if not config['smtp_password']:
                config['smtp_password'] = (email_cfg.get('smtp_password', '') or '').strip()

            if not config['sender']:
                config['sender'] = (email_cfg.get('sender', '') or '').strip()

            if os.getenv('EMAIL_USE_SSL') is None:
                config['use_ssl'] = bool(email_cfg.get('use_ssl', config['use_ssl']))

            if not config['recipients']:
                file_recipients = email_cfg.get('recipients', [])
                if isinstance(file_recipients, str):
                    config['recipients'] = [r.strip() for r in file_recipients.split(',') if r.strip()]
                elif isinstance(file_recipients, list):
                    config['recipients'] = [str(r).strip() for r in file_recipients if str(r).strip()]
    except Exception:
        pass

    if not config['sender']:
        config['sender'] = config['smtp_user']

    return config


def run_daily_report(app, db_manager):
    """Generate and email the daily report"""
    try:
        logger.info("Starting daily report generation...")
        filepath = generate_daily_report(app, db_manager)
        if filepath:
            logger.info(f"Report generated successfully: {filepath}")
            email_config = get_email_config()
            
            # Check email configuration
            if not email_config.get('smtp_host'):
                logger.warning("Email SMTP host not configured - report saved locally only")
            elif not email_config.get('recipients'):
                logger.warning("Email recipients not configured - report saved locally only")
            else:
                logger.info(f"Sending email to {email_config.get('recipients')}")
                success = send_report_email(filepath, email_config)
                if success:
                    logger.info("Daily report email sent successfully")
                else:
                    logger.warning("Failed to send daily report email (see logs for details)")
            return True
        else:
            logger.warning("No active stores found or report generation failed")
            return False
    except Exception as e:
        logger.error(f"Daily report failed: {e}", exc_info=True)
        return False


def run_monthly_report(app, db_manager, target_month=None):
    """Generate and email the monthly report for a month (defaults to previous month)."""
    try:
        if target_month:
            year, month = _parse_month_input(target_month)
            report_start, _ = _get_month_range(year, month)
        else:
            report_start, _ = _get_previous_month_range()

        month_label = report_start.strftime('%B %Y')
        month_key = report_start.strftime('%Y-%m')
        logger.info(f"Starting monthly report generation for {month_label}...")

        filepath = generate_monthly_report(app, db_manager, target_month=month_key)
        if filepath:
            logger.info(f"Monthly report generated successfully: {filepath}")
            email_config = get_email_config()

            if not email_config.get('smtp_host'):
                logger.warning("Email SMTP host not configured - monthly report saved locally only")
            elif not email_config.get('recipients'):
                logger.warning("Email recipients not configured - monthly report saved locally only")
            else:
                subject = f"Monthly AI Alerts Report - {month_label}"
                body = f"""Dear Team,

Please find attached the Monthly AI Alerts Report for {month_label}.

This is an automated report generated by Sakshi.AI Video Analytics Platform.

Regards,
Sakshi.AI System"""
                logger.info(f"Sending monthly email to {email_config.get('recipients')}")
                success = send_report_email(filepath, email_config, report_subject=subject, report_body=body)
                if success:
                    logger.info("Monthly report email sent successfully")
                else:
                    logger.warning("Failed to send monthly report email (see logs for details)")
        else:
            logger.warning("No active stores found or monthly report generation failed")
    except Exception as e:
        logger.error(f"Monthly report failed: {e}", exc_info=True)


def start_daily_report_scheduler(app, db_manager):
    """Start a background thread that sends the daily report at 10:30 AM IST"""

    def scheduler_loop():
        logger.info("Daily report scheduler started (10:30 AM IST)")
        while True:
            try:
                now = datetime.now(IST)
                target_today = now.replace(hour=10, minute=30, second=0, microsecond=0)

                # If scheduler starts/restarts after 10:30, perform a catch-up run
                # for yesterday if that day's report file is missing.
                if now >= target_today:
                    report_date = (now - timedelta(days=1)).strftime('%Y-%m-%d')
                    if _daily_report_exists_for_date(report_date):
                        logger.info(
                            f"Daily report for {report_date} already exists; scheduling next run for tomorrow"
                        )
                        target = target_today + timedelta(days=1)
                    else:
                        logger.warning(
                            f"Daily report for {report_date} missing after scheduled time; running catch-up now"
                        )
                        success = run_daily_report(app, db_manager)
                        if success or _daily_report_exists_for_date(report_date):
                            target = target_today + timedelta(days=1)
                        else:
                            # Retry in 30 minutes if generation fails after scheduled time.
                            target = now + timedelta(minutes=30)
                            logger.warning(
                                "Catch-up daily report run did not produce a report file; retrying in 30 minutes"
                            )
                else:
                    target = target_today

                wait_seconds = max((target - now).total_seconds(), 1)
                logger.info(f"Next daily report scheduled at {target.strftime('%Y-%m-%d %H:%M:%S')} IST "
                          f"(in {wait_seconds/3600:.1f} hours)")
                time.sleep(wait_seconds)

                # Run only on the regular schedule wake-up.
                if target.hour == 10 and target.minute == 30:
                    run_daily_report(app, db_manager)

            except Exception as e:
                logger.error(f"Daily report scheduler error: {e}", exc_info=True)
                # Sleep 60s on error before retrying
                time.sleep(60)

    thread = threading.Thread(target=scheduler_loop, daemon=True, name='DailyReportScheduler')
    thread.start()
    return thread


def start_monthly_report_scheduler(app, db_manager):
    """Start a background thread that sends monthly report at 10:00 AM IST on day 1."""

    def scheduler_loop():
        logger.info("Monthly report scheduler started (1st day, 10:00 AM IST)")
        while True:
            try:
                now = datetime.now(IST)
                target = now.replace(day=1, hour=10, minute=0, second=0, microsecond=0)
                if now >= target:
                    if target.month == 12:
                        target = target.replace(year=target.year + 1, month=1)
                    else:
                        target = target.replace(month=target.month + 1)

                wait_seconds = (target - now).total_seconds()
                logger.info(
                    f"Next monthly report scheduled at {target.strftime('%Y-%m-%d %H:%M:%S')} IST "
                    f"(in {wait_seconds / 3600:.1f} hours)"
                )
                time.sleep(wait_seconds)
                run_monthly_report(app, db_manager)

            except Exception as e:
                logger.error(f"Monthly report scheduler error: {e}", exc_info=True)
                time.sleep(60)

    thread = threading.Thread(target=scheduler_loop, daemon=True, name='MonthlyReportScheduler')
    thread.start()
    return thread
